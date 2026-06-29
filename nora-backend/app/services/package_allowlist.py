"""Package allowlist (XLSX → JSON) and PEP 440 version matching."""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal

from packaging.specifiers import SpecifierSet
from packaging.utils import canonicalize_name
from packaging.version import Version

from app.services.data_bootstrap import resolve_data_root

SCHEMA = "nora.package-allowlist/1"
_ALLOWLIST_REL = "noraops/packages/allowlist.json"

ImportRowStatus = Literal["allowed", "review", "skipped"]

_REVIEW_VERSION_HINTS = re.compile(
    r"(要確認|未定|不明|TBD|N/A|latest|master|main|dev|snapshot|随時|検討|保留)",
    re.IGNORECASE,
)
_JA_CHARS = re.compile(r"[\u3040-\u30ff\u4e00-\u9fff]")
_WILDCARD_VERSION = re.compile(r"^[\d*]+(?:\.[\d*]+)*$")
_PKG_NAME_OK = re.compile(r"^[a-z0-9]([a-z0-9._-]*[a-z0-9])?$")


@dataclass
class XlsxImportConfig:
    header_row: int = 9
    category_col: str = "A"
    category_filter: str = "Pythonライブラリ"
    package_col: str = "B"
    version_col: str = "C"
    maintenance_date_col: str = "D"

    def normalized(self) -> XlsxImportConfig:
        return XlsxImportConfig(
            header_row=max(1, int(self.header_row)),
            category_col=_normalize_col(self.category_col),
            category_filter=str(self.category_filter or "").strip(),
            package_col=_normalize_col(self.package_col) or "B",
            version_col=_normalize_col(self.version_col) or "C",
            maintenance_date_col=_normalize_col(self.maintenance_date_col),
        )

    def to_dict(self) -> dict[str, Any]:
        cfg = self.normalized()
        return {
            "headerRow": cfg.header_row,
            "categoryCol": cfg.category_col or None,
            "categoryFilter": cfg.category_filter or None,
            "packageCol": cfg.package_col,
            "versionCol": cfg.version_col,
            "maintenanceDateCol": cfg.maintenance_date_col or None,
        }


def allowlist_path() -> Path:
    return resolve_data_root() / _ALLOWLIST_REL


def _empty_document() -> dict:
    now = datetime.now(timezone.utc).isoformat()
    return {"schema": SCHEMA, "updatedAt": now, "packages": []}


def load_allowlist() -> dict:
    path = allowlist_path()
    if not path.is_file():
        return _empty_document()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return _empty_document()
    if not isinstance(data, dict):
        return _empty_document()
    data.setdefault("schema", SCHEMA)
    data.setdefault("packages", [])
    return data


def save_allowlist(data: dict) -> dict:
    path = allowlist_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    data = dict(data)
    data["schema"] = SCHEMA
    data["updatedAt"] = datetime.now(timezone.utc).isoformat()
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return data


def compute_etag(data: dict | None = None) -> str:
    payload = data if data is not None else load_allowlist()
    raw = json.dumps(payload, sort_keys=True, ensure_ascii=False).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()[:16]


def normalize_pkg_name(name: str) -> str:
    try:
        return canonicalize_name(str(name or "").strip())
    except Exception:
        return str(name or "").strip().lower().replace("_", "-")


def _normalize_col(col: str | None) -> str:
    return str(col or "").strip().upper()


def column_letter_to_index(col: str) -> int:
    """Excel column letter (A, B, …) → 0-based index."""
    s = _normalize_col(col)
    if not s:
        raise ValueError("列番号が空です")
    if not re.fullmatch(r"[A-Z]+", s):
        raise ValueError(f"列番号が不正です: {col}")
    idx = 0
    for ch in s:
        idx = idx * 26 + (ord(ch) - ord("A") + 1)
    return idx - 1


def _cell_value(row: tuple[Any, ...] | None, col: str) -> Any:
    if not row or not col:
        return None
    idx = column_letter_to_index(col)
    if idx >= len(row):
        return None
    return row[idx]


def _wildcard_version_to_spec(s: str) -> str | None:
    parts = s.split(".")
    if "*" not in parts:
        return None
    star_i = parts.index("*")
    if star_i == 0:
        return ""
    prefix = parts[:star_i]
    if not all(re.fullmatch(r"\d+", p) for p in prefix):
        return None
    lower: list[str] = []
    for i, part in enumerate(parts):
        if part == "*":
            lower.append("0")
            break
        lower.append(part)
    while len(lower) < len(parts):
        lower.append("0")
    lower_ver = ".".join(lower)
    upper_prefix = prefix[:-1] + [str(int(prefix[-1]) + 1)]
    upper_ver = ".".join(upper_prefix)
    return f">={lower_ver},<{upper_ver}"


def _validate_specifier(spec: str) -> bool:
    if not spec:
        return True
    try:
        SpecifierSet(spec)
        return True
    except Exception:
        return False


def _convert_version_spec(s: str) -> str | None:
    if _WILDCARD_VERSION.fullmatch(s) and "*" in s:
        if s.strip() == "*":
            return ""
        return _wildcard_version_to_spec(s)
    if re.fullmatch(r"\d+\.\*", s):
        major = s.split(".", 1)[0]
        return f">={major}.0,<{int(major) + 1}.0"
    if re.fullmatch(r"\d+\.\d+\.\*", s):
        parts = s.split(".")
        return f">={parts[0]}.{parts[1]}.0,<{parts[0]}.{int(parts[1]) + 1}.0"
    try:
        SpecifierSet(s)
        return s
    except Exception:
        pass
    try:
        Version(s)
        return f"=={s}"
    except Exception:
        return None


def _looks_suspicious_version_text(s: str) -> bool:
    if len(s) > 80 or "\n" in s or "http" in s.lower():
        return True
    if _REVIEW_VERSION_HINTS.search(s):
        return True
    if _JA_CHARS.search(s) and s not in ("すべて", "全バージョン"):
        return True
    if re.search(r"\s+or\s+|および|／|/", s, re.IGNORECASE):
        return True
    return False


def _looks_suspicious_package_name(name: str) -> str | None:
    raw = str(name or "").strip()
    if not raw:
        return "パッケージ名が空です"
    if raw.startswith("#"):
        return "コメント行"
    if " " in raw or "://" in raw or "@" in raw:
        return "パッケージ名の形式が不正です"
    key = normalize_pkg_name(raw)
    if not key or not _PKG_NAME_OK.fullmatch(key):
        return "PyPI パッケージ名として認識できません"
    return None


def parse_version_for_import(raw: str) -> tuple[ImportRowStatus, str, str | None]:
    """Return (status, versionSpec, reason). status is allowed or review."""
    s = str(raw or "").strip()
    if not s or s in ("-", "—", "ー", "*", "any", "すべて", "全バージョン"):
        return "allowed", "", None
    if _looks_suspicious_version_text(s):
        return "review", "", f"バージョン要確認: {s}"
    converted = _convert_version_spec(s)
    if converted is None:
        return "review", "", f"バージョン制約を解釈できません: {s}"
    if not _validate_specifier(converted):
        return "review", "", f"PEP 440 として無効: {s}"
    return "allowed", converted, None


def _parse_version_spec(raw: str) -> str:
    status, spec, _ = parse_version_for_import(raw)
    if status == "review":
        return raw.strip()
    return spec


def version_matches(spec_raw: str, version: str | None) -> bool:
    spec_text = _parse_version_spec(spec_raw)
    if not spec_text:
        return True
    if not version:
        return True
    try:
        return Version(version) in SpecifierSet(spec_text)
    except Exception:
        return False


def _entries_by_name(data: dict) -> dict[str, list[dict]]:
    out: dict[str, list[dict]] = {}
    for row in data.get("packages") or []:
        if not isinstance(row, dict):
            continue
        name = normalize_pkg_name(str(row.get("name") or ""))
        if not name:
            continue
        out.setdefault(name, []).append(row)
    return out


def is_package_allowed(name: str, version: str | None = None, data: dict | None = None) -> bool:
    data = data or load_allowlist()
    key = normalize_pkg_name(name)
    entries = _entries_by_name(data).get(key)
    if not entries:
        return False
    for ent in entries:
        spec = str(ent.get("versionSpec") or ent.get("version_spec") or "")
        if version_matches(spec, version):
            return True
    return False


def classify_packages(
    packages: list[dict[str, str]],
    data: dict | None = None,
) -> dict:
    """packages: [{name, version?}] → approved / unapproved lists."""
    data = data or load_allowlist()
    approved: list[dict] = []
    unapproved: list[dict] = []
    for item in packages:
        name = str(item.get("name") or item.get("package") or "").strip()
        if not name:
            continue
        version = (item.get("version") or "").strip() or None
        row = {"name": normalize_pkg_name(name), "version": version or ""}
        if is_package_allowed(name, version, data):
            approved.append(row)
        else:
            unapproved.append(row)
    return {"approved": approved, "unapproved": unapproved}


def _import_row_dict(
    *,
    row_num: int,
    status: ImportRowStatus,
    name: str = "",
    version_raw: str = "",
    version_spec: str = "",
    reason: str | None = None,
    category: str | None = None,
    maintenance_date: str | None = None,
) -> dict[str, Any]:
    return {
        "row": row_num,
        "status": status,
        "name": name,
        "versionRaw": version_raw,
        "versionSpec": version_spec,
        "reason": reason,
        "category": category,
        "maintenanceDate": maintenance_date,
    }


def parse_xlsx_import(content: bytes, config: XlsxImportConfig | None = None) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("openpyxl is required for XLSX upload") from e

    from io import BytesIO

    cfg = (config or XlsxImportConfig()).normalized()
    if not cfg.package_col or not cfg.version_col:
        raise ValueError("パッケージ名列・バージョン列は必須です")

    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    allowed: list[dict[str, Any]] = []
    review: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    packages: list[dict[str, Any]] = []
    seen: set[str] = set()
    data_start = cfg.header_row + 1

    for row_num, row in enumerate(
        ws.iter_rows(min_row=data_start, values_only=True),
        start=data_start,
    ):
        if not row:
            continue

        category_val = None
        if cfg.category_col:
            category_val = _cell_value(row, cfg.category_col)
            category_text = str(category_val or "").strip()
            if cfg.category_filter:
                if category_text != cfg.category_filter:
                    skipped.append(
                        _import_row_dict(
                            row_num=row_num,
                            status="skipped",
                            category=category_text or None,
                            reason=f"分類不一致（期待: {cfg.category_filter}）",
                        )
                    )
                    continue

        pkg_raw = _cell_value(row, cfg.package_col)
        ver_raw = _cell_value(row, cfg.version_col)
        maint_raw = (
            _cell_value(row, cfg.maintenance_date_col)
            if cfg.maintenance_date_col
            else None
        )

        name_raw = str(pkg_raw or "").strip()
        version_raw = str(ver_raw or "").strip()
        maintenance_date = str(maint_raw or "").strip() or None

        pkg_issue = _looks_suspicious_package_name(name_raw)
        if pkg_issue == "コメント行" or (not name_raw and not version_raw):
            continue
        if pkg_issue:
            review.append(
                _import_row_dict(
                    row_num=row_num,
                    status="review",
                    name=name_raw,
                    version_raw=version_raw,
                    reason=pkg_issue,
                    category=str(category_val or "").strip() or None,
                    maintenance_date=maintenance_date,
                )
            )
            continue

        key = normalize_pkg_name(name_raw)
        if key in seen:
            skipped.append(
                _import_row_dict(
                    row_num=row_num,
                    status="skipped",
                    name=key,
                    version_raw=version_raw,
                    reason="重複（先の行を採用）",
                    category=str(category_val or "").strip() or None,
                    maintenance_date=maintenance_date,
                )
            )
            continue

        ver_status, version_spec, ver_reason = parse_version_for_import(version_raw)
        if ver_status == "review":
            review.append(
                _import_row_dict(
                    row_num=row_num,
                    status="review",
                    name=key,
                    version_raw=version_raw,
                    reason=ver_reason,
                    category=str(category_val or "").strip() or None,
                    maintenance_date=maintenance_date,
                )
            )
            continue

        seen.add(key)
        entry: dict[str, Any] = {"name": key, "versionSpec": version_spec}
        if maintenance_date:
            entry["maintenanceDate"] = maintenance_date
        if category_val and str(category_val).strip():
            entry["category"] = str(category_val).strip()
        packages.append(entry)
        allowed.append(
            _import_row_dict(
                row_num=row_num,
                status="allowed",
                name=key,
                version_raw=version_raw,
                version_spec=version_spec,
                category=str(category_val or "").strip() or None,
                maintenance_date=maintenance_date,
            )
        )

    wb.close()
    return {
        "config": cfg.to_dict(),
        "allowed": allowed,
        "review": review,
        "skipped": skipped,
        "packages": packages,
    }


def _package_snapshot(packages: list[dict]) -> dict[str, dict[str, str]]:
    out: dict[str, dict[str, str]] = {}
    for row in packages:
        if not isinstance(row, dict):
            continue
        name = normalize_pkg_name(str(row.get("name") or ""))
        if not name:
            continue
        out[name] = {
            "name": name,
            "versionSpec": str(row.get("versionSpec") or row.get("version_spec") or ""),
            "maintenanceDate": str(row.get("maintenanceDate") or ""),
            "category": str(row.get("category") or ""),
        }
    return out


def compute_allowlist_diff(
    old_packages: list[dict],
    new_packages: list[dict],
) -> dict[str, Any]:
    """Compare previous and new allowlist entries (by package name)."""
    old_map = _package_snapshot(old_packages)
    new_map = _package_snapshot(new_packages)
    old_names = set(old_map)
    new_names = set(new_map)

    added: list[dict[str, Any]] = []
    removed: list[dict[str, Any]] = []
    changed: list[dict[str, Any]] = []
    unchanged_count = 0

    for name in sorted(new_names - old_names):
        added.append({**new_map[name], "change": "added"})

    for name in sorted(old_names - new_names):
        removed.append({**old_map[name], "change": "removed"})

    for name in sorted(old_names & new_names):
        before = old_map[name]
        after = new_map[name]
        fields: dict[str, dict[str, str]] = {}
        for field in ("versionSpec", "maintenanceDate", "category"):
            if before.get(field) != after.get(field):
                fields[field] = {"from": before.get(field, ""), "to": after.get(field, "")}
        if fields:
            changed.append({"name": name, "fields": fields, "before": before, "after": after})
        else:
            unchanged_count += 1

    return {
        "added": added,
        "removed": removed,
        "changed": changed,
        "unchangedCount": unchanged_count,
        "beforeCount": len(old_map),
        "afterCount": len(new_map),
    }


def ingest_xlsx(content: bytes, config: XlsxImportConfig | None = None) -> dict:
    result = parse_xlsx_import(content, config)
    data = load_allowlist()
    previous_packages = list(data.get("packages") or [])
    diff = compute_allowlist_diff(previous_packages, result["packages"])
    data["packages"] = result["packages"]
    data["source"] = "xlsx"
    data["lastImport"] = {
        "at": datetime.now(timezone.utc).isoformat(),
        "config": result["config"],
        "counts": {
            "allowed": len(result["allowed"]),
            "review": len(result["review"]),
            "skipped": len(result["skipped"]),
        },
        "review": result["review"],
        "skipped": result["skipped"],
        "diff": diff,
    }
    save_allowlist(data)
    return {**result, "allowlist": data, "diff": diff}


def allowlist_summary(data: dict | None = None) -> dict:
    data = data or load_allowlist()
    packages = data.get("packages") or []
    last = data.get("lastImport") or {}
    return {
        "schema": data.get("schema", SCHEMA),
        "updatedAt": data.get("updatedAt"),
        "packageCount": len(packages),
        "etag": compute_etag(data),
        "lastImport": last,
    }
